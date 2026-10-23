"""Generate non-sensitive document fixtures inside the isolated Docling runtime."""

from pathlib import Path
import sys


def text_pdf(path):
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R 5 0 R] /Count 2 >>"]
    for label in ["Premiere page ATLAS", "Deuxieme page BOREAL"]:
        stream = f"BT /F1 22 Tf 60 730 Td ({label}) Tj ET".encode()
        content_id = len(objects) + 2
        objects.extend([
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 7 0 R >> >> /Contents {content_id} 0 R >>".encode(),
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream",
        ])
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(data)


def generate(root):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from PIL import Image, ImageDraw, ImageFont

    root.mkdir(parents=True, exist_ok=True)
    text_pdf(root / "native.pdf")
    document = Document()
    document.add_heading("École et café", 0)
    document.add_paragraph("Document témoin ORION : caractères accentués conservés.")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("Article", "Prix"), ("Stylo", "3,50 euros")]):
        for cell, text in zip(row.cells, values):
            cell.text = text
    document.save(root / "office.docx")
    workbook = Workbook()
    workbook.active.append(["Article", "Quantité"])
    workbook.active.append(["Cellule VEGA", 17])
    workbook.save(root / "table.xlsx")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Diapositive ALTAIR"
    slide.placeholders[1].text = "École et café"
    presentation.save(root / "slides.pptx")
    (root / "page.html").write_text("<html><body><h1>École et café</h1><p>Texte SIRIUS</p></body></html>", encoding="utf-8")
    image = Image.new("RGB", (1600, 1000), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 48)
    for y, line in enumerate(["Document local", "Reference AZUR 7319", "École et café", "Montant 42 euros"]):
        draw.text((100, 100 + y * 130), line, fill="black", font=font)
    image.save(root / "scan.png")
    image.save(root / "scan.pdf", resolution=150)
    (root / "broken.pdf").write_bytes(b"this is not a PDF")


if __name__ == "__main__":
    generate(Path(sys.argv[1]))
